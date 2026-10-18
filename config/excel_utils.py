import openpyxl
from datetime import date, datetime
from numbers import Number
from django.utils import timezone
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse


def export_queryset_to_excel(filename, sheet_title, headers, rows_data):


    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]


    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="B80003", end_color="B80003", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")


    ws['A1'] = f"Ilustre Municipalidad de La Serena • {sheet_title}"
    ws['A1'].font = Font(name="Arial", size=13, bold=True, color="B80003")
    ws.row_dimensions[1].height = 25


    ws.row_dimensions[3].height = 22
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin


    for row_idx, row_values in enumerate(rows_data, 4):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, value in enumerate(row_values, 1):
            if isinstance(value, datetime) and timezone.is_aware(value):
                value = timezone.localtime(value).replace(tzinfo=None)
            native = value if isinstance(value, (Number, date, datetime)) else str(value) if value is not None else ""
            cell = ws.cell(row=row_idx, column=col_idx, value=native)
            if isinstance(native, str):

                cell.data_type = 's'
            cell.font = Font(name="Arial", size=10)
            cell.border = border_thin
            cell.alignment = align_center if col_idx in [1, len(row_values)] else align_left


    for col_idx, header_text in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header_text)
        for row_values in rows_data:
            if col_idx - 1 < len(row_values):
                val_str = str(row_values[col_idx - 1] or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(ws.max_row, 3)}'

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response
