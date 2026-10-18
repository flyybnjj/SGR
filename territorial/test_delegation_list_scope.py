from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from usuarios.models import Rol, Usuario
from .models import DelegacionMunicipal


class DelegationListScopeTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.own = DelegacionMunicipal.objects.create(nombre='Sede autorizada', direccion='Principal')
        cls.foreign = DelegacionMunicipal.objects.create(nombre='Sede ajena privada', direccion='Otra calle')
        role = Rol.objects.create(nombre='Gestor', puede_gestionar=True)
        cls.manager = Usuario.objects.create_user(username='manager', rol=role, delegacion=cls.own)
        cls.root = Usuario.objects.create_superuser(username='root', password='unused')

    def test_manager_list_search_and_excel_only_include_assigned_seat(self):
        self.client.force_login(self.manager)
        for query in ('', '?q=Sede', '?q=ajena'):
            with self.subTest(query=query):
                page = self.client.get(reverse('territorial:delegaciones_list') + query)
                self.assertEqual(page.status_code, 200)
                self.assertEqual(list(page.context['page_obj'].object_list), [] if 'ajena' in query else [self.own])
                export = self.client.get(reverse('territorial:delegaciones_excel') + query)
                rows = list(load_workbook(BytesIO(export.content)).active.values)
                self.assertFalse(any(self.foreign.nombre in row for row in rows))

    def test_root_can_list_and_export_all_seats(self):
        self.client.force_login(self.root)
        page = self.client.get(reverse('territorial:delegaciones_list'))
        self.assertCountEqual(page.context['page_obj'].object_list, [self.own, self.foreign])
        export = self.client.get(reverse('territorial:delegaciones_excel'))
        rows = list(load_workbook(BytesIO(export.content)).active.values)
        self.assertTrue(any(self.foreign.nombre in row for row in rows))
