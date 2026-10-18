"""Regresiones de auditoría; ejecutar exclusivamente con settings aislados."""
from io import BytesIO

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from atenciones.models import Accion, Atencion, MetaPeriodo, TipoAtencion
from territorial.forms import DelegacionMunicipalForm
from territorial.models import DelegacionMunicipal, ItemMeta
from usuarios.models import Rol, Usuario
from vecinos.models import Vecino


class TerritorialLogicAuditTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.role = Rol.objects.create(nombre='Auditoría territorial', puede_gestionar=True)
        cls.seat = DelegacionMunicipal.objects.create(nombre='Sede de auditoría', direccion='Calle Uno')
        cls.staff = Usuario.objects.create_user(username='audit-territorial-staff', rol=cls.role, delegacion=cls.seat)
        cls.root = Usuario.objects.create_superuser(username='audit-territorial-root')

    def test_inactivated_workplace_does_not_keep_operational_dashboard_data(self):
        person = Vecino.objects.create(nombre='Vecina de auditoría', territorio=self.seat)
        kind = TipoAtencion.objects.create(nombre='Atención de auditoría')
        Atencion.objects.create(vecino=person, delegacion=self.seat, funcionario=self.staff, tipo_atencion=kind, descripcion='Consulta')
        today = timezone.localdate()
        meta = MetaPeriodo.objects.create(delegacion=self.seat, item_meta=ItemMeta.objects.create(nombre='Meta de auditoría'), objetivo=1, inicio=today, fin=today)
        Accion.objects.create(titulo='Pendiente privada', descripcion='Detalle', delegacion=self.seat, creador=self.staff, meta=meta)
        self.client.force_login(self.staff)
        self.seat.activo = False
        self.seat.save(update_fields=['activo'])
        response = self.client.get(reverse('dashboard:inicio'))
        if response.status_code == 403:
            return
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total_atenciones'], 0)
        self.assertEqual(response.context['personas_atendidas'], 0)
        self.assertEqual(response.context['acciones_pendientes'], 0)
        self.assertEqual(list(response.context['metas_actuales']), [])
        self.assertEqual(list(response.context['pendientes_recientes']), [])

    def test_delegation_excel_uses_effective_unconfirmed_color(self):
        self.seat.color = '#1A33EA'
        self.seat.color_configurado = False
        self.seat.save(update_fields=['color', 'color_configurado'])
        self.client.force_login(self.root)
        response = self.client.get(reverse('territorial:delegaciones_excel'), {'q': self.seat.nombre})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.cell(row=4, column=7).value, '#B80003')

    def test_interleaved_assignments_cannot_leave_incoherent_responsibility(self):
        delegate = Usuario.objects.create_user(username='audit-free-delegate', rol=self.role)
        forms = [DelegacionMunicipalForm(data={
            'nombre': name, 'direccion': 'Calle Dos', 'comuna': 'La Serena',
            'delegado': delegate.pk, 'activo': 'on',
        }) for name in ('Primera sede simultánea', 'Segunda sede simultánea')]
        for form in forms:
            self.assertTrue(form.is_valid(), form.errors)
        forms[0].save()
        try:
            forms[1].save()
        except ValidationError:
            pass
        delegate.refresh_from_db()
        self.assertFalse(DelegacionMunicipal.objects.filter(delegado=delegate).exclude(pk=delegate.delegacion_id).exists())

    def test_legacy_unassigned_delegate_cannot_take_a_second_seat(self):
        delegate = Usuario.objects.create_user(username='audit-legacy-delegate', rol=self.role)
        DelegacionMunicipal.objects.create(nombre='Responsabilidad heredada', direccion='Calle Tres', delegado=delegate)
        form = DelegacionMunicipalForm(data={
            'nombre': 'Nueva responsabilidad incompatible', 'direccion': 'Calle Cuatro',
            'comuna': 'La Serena', 'delegado': delegate.pk, 'activo': 'on',
        })
        self.assertFalse(form.is_valid())
        self.assertIn('delegado', form.errors)
